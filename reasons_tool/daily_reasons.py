# -*- coding: utf-8 -*-
"""
상승 이유 자동 수집기 (PC에서 Claude 없이 실행)

하는 일
  1) 네이버 검색 Open API(뉴스)로 '특징주' 기사를 조회합니다. (웹 크롤링 아님)
  2) 기사 제목에서 종목명과 이유를 뽑아 한 줄로 만듭니다.
  3) data/reasons.json(대시보드용)과 상승이유_누적.xlsx(엑셀 누적)에 추가합니다. 같은 기사 링크는 중복 제거.
  4) (선택) GitHub 저장소로 자동 업로드합니다.

필요한 것
  - Python 3.8 이상 (표준 라이브러리만 사용. 엑셀 누적에만 openpyxl이 있으면 쓰고, 없으면 건너뜁니다)
  - 네이버 개발자센터에서 받은 Client ID / Client Secret  → config.ini
  - 대시보드 저장소를 내려받은 폴더 (data/tick.json 이 있어야 종목명을 찾습니다)

실행:  python daily_reasons.py            (config.ini 는 같은 폴더)
"""
import os, sys, re, json, html, time, subprocess, datetime, configparser
import urllib.request, urllib.parse, urllib.error
from email.utils import parsedate_to_datetime
from urllib.parse import urlparse

HERE = os.path.dirname(os.path.abspath(__file__))
KST = datetime.timezone(datetime.timedelta(hours=9))
LOG = os.path.join(HERE, 'log.txt')


def log(msg):
    line = f'[{datetime.datetime.now(KST):%Y-%m-%d %H:%M:%S}] {msg}'
    print(line)
    try:
        with open(LOG, 'a', encoding='utf-8') as f:
            f.write(line + '\n')
    except OSError:
        pass


# ------------------------------------------------------------------ 설정
def load_cfg():
    p = os.path.join(HERE, 'config.ini')
    if not os.path.exists(p):
        sys.exit('config.ini 가 없습니다. config.ini.sample 을 복사해서 값을 채워 주세요.')
    c = configparser.ConfigParser()
    c.read(p, encoding='utf-8-sig')
    g = lambda s, k, d='': c.get(s, k, fallback=d).strip()
    cfg = dict(
        cid=g('naver', 'client_id'), sec=g('naver', 'client_secret'),
        repo=g('paths', 'repo'),
        queries=[q.strip() for q in g('options', 'queries', '특징주').split(',') if q.strip()],
        days_keep=int(g('options', 'days_keep', '90')),
        fetch_days=int(g('options', 'fetch_days', '3')),
        push=g('options', 'push', 'no').lower() in ('yes', 'y', 'true', '1'),
    )
    if not cfg['cid'] or not cfg['sec'] or 'YOUR' in cfg['cid'].upper():
        sys.exit('config.ini 의 client_id / client_secret 을 채워 주세요.')
    if not cfg['repo'] or not os.path.isdir(os.path.join(cfg['repo'], 'data')):
        sys.exit('config.ini 의 repo 경로가 잘못되었습니다. 대시보드 저장소 폴더(안에 data 폴더가 있는 곳)를 적어 주세요.')
    return cfg


# ------------------------------------------------------------------ 네이버 뉴스 API
API = 'https://openapi.naver.com/v1/search/news.json'


def naver_get(cfg, query, start, display=100, retries=4):
    url = f'{API}?query={urllib.parse.quote(query)}&display={display}&start={start}&sort=date'
    req = urllib.request.Request(url, headers={'X-Naver-Client-Id': cfg['cid'], 'X-Naver-Client-Secret': cfg['sec']})
    for a in range(retries):
        try:
            with urllib.request.urlopen(req, timeout=20) as r:
                return json.loads(r.read().decode('utf-8'))
        except urllib.error.HTTPError as e:
            body = e.read().decode('utf-8', 'ignore')[:200]
            if e.code in (401, 403):
                sys.exit(f'네이버 인증 오류({e.code}). Client ID/Secret 과 "검색" API 사용 설정을 확인하세요. {body}')
            if e.code == 429:           # 호출 한도 초과 -> 잠시 후 재시도
                time.sleep(2 * (a + 1)); continue
            log(f'HTTP {e.code}: {body}'); time.sleep(1.5 * (a + 1))
        except (urllib.error.URLError, TimeoutError, ValueError) as e:
            log(f'네트워크 오류: {e}'); time.sleep(1.5 * (a + 1))
    return None


def fetch_items(cfg, now=None, getter=None):
    getter = getter or naver_get
    now = now or datetime.datetime.now(KST)
    cutoff = now - datetime.timedelta(days=cfg['fetch_days'])
    out = []
    for q in cfg['queries']:
        for start in range(1, 1001, 100):      # API는 최대 1000건까지
            j = getter(cfg, q, start)
            if not j or not j.get('items'):
                break
            oldest = None
            for it in j['items']:
                try:
                    dt = parsedate_to_datetime(it['pubDate']).astimezone(KST)
                except Exception:
                    continue
                oldest = dt if oldest is None or dt < oldest else oldest
                if dt >= cutoff:
                    out.append((it, dt))
            if oldest and oldest < cutoff:
                break
            time.sleep(0.12)
    return out


# ------------------------------------------------------------------ 제목 가공
TAG = re.compile(r'<[^>]+>')
LEAD = re.compile(r'^\s*[\[\(<【]?\s*(특징주|시황|속보|종합)\s*[\]\)>】]?\s*[:\-·]?\s*')
UPW = ('급등', '강세', '상승', '신고가', '상한가', '폭등', '껑충', '급반등', '반등', '오름', '↑', '들썩', '치솟', '뛰')
DNW = ('급락', '약세', '하락', '하한가', '폭락', '내림', '↓', '밀려', '추락', '털썩')
HANGUL = re.compile(r'[가-힣A-Za-z0-9]')


def clean_title(t):
    t = html.unescape(TAG.sub('', t or '')).replace('\u200b', '').strip()
    return re.sub(r'\s+', ' ', t)


def direction(t):
    pos = [(t.find(w), 'up') for w in UPW if w in t] + [(t.find(w), 'down') for w in DNW if w in t]
    return min(pos)[1] if pos else '?'


def find_names(title, names_by_len):
    """제목 안의 종목명을 찾는다. 긴 이름 우선, 겹치지 않게. 앞 글자가 한글/영문/숫자면 건너뜀."""
    found, taken = [], [False] * len(title)
    for nm in names_by_len:
        start = 0
        while True:
            i = title.find(nm, start)
            if i < 0:
                break
            j = i + len(nm)
            ok = not any(taken[i:j])
            if ok and i > 0 and HANGUL.match(title[i - 1]):
                ok = False
            if ok and len(nm) <= 2:        # 두 글자 이름은 오탐이 많아서 뒤에 한글이 바로 이어지면 제외
                if j < len(title) and re.match(r'[가-힣]', title[j]):
                    ok = False
            if ok:
                found.append((i, nm))
                for k in range(i, j): taken[k] = True
                break
            start = i + 1
    found.sort()
    return [n for _, n in found]


def reason_of(title, name, all_names=()):
    t = LEAD.sub('', title)
    t = re.sub(r'^\s*\[[^\]]{1,12}\]\s*', '', t)          # 맨 앞 [태그] 제거
    for n in sorted(set(all_names) | {name}, key=len, reverse=True):   # 같은 기사에 나온 다른 종목명도 제거
        t = t.replace(n, '', 1)
    t = re.sub(r'^[\s·,]+', '', t)
    t = re.sub(r'^[\s,·…\-–—:\'\"“”‘’()]+', '', t).strip()
    return t[:120]


def build_rows(items, names_by_len):
    rows = []
    for it, dt in items:
        title = clean_title(it.get('title'))
        if not title:
            continue
        url = it.get('originallink') or it.get('link') or ''
        src = urlparse(url).netloc.replace('www.', '')
        found = find_names(title, names_by_len)
        for nm in found:
            rows.append(dict(d=f'{dt:%Y-%m-%d}', t=f'{dt:%H:%M}', n=nm, r=reason_of(title, nm, found), h=title, u=url, s=src, dir=direction(title)))
    return rows


# ------------------------------------------------------------------ 누적 저장
def load_names(repo):
    with open(os.path.join(repo, 'data', 'tick.json'), encoding='utf-8') as f:
        names = {t[1] for t in json.load(f) if len(t) > 1 and len(t[1]) >= 2}
    return sorted(names, key=lambda x: (-len(x), x))


def merge(old, new, days_keep, today):
    seen = {(x['u'], x['n']) for x in old}
    add = [x for x in new if (x['u'], x['n']) not in seen]
    allr = old + add
    cut = f'{today - datetime.timedelta(days=days_keep):%Y-%m-%d}'
    allr = [x for x in allr if x['d'] >= cut]
    allr.sort(key=lambda x: (x['d'], x['t']), reverse=True)
    return allr, add


def save_json(path, rows):
    tmp = path + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(rows, f, ensure_ascii=False, separators=(',', ':'))
    os.replace(tmp, path)


def save_excel(path, add):
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        log('openpyxl 이 없어 엑셀 누적은 건너뜁니다. (pip install openpyxl)')
        return
    if os.path.exists(path):
        wb = load_workbook(path); ws = wb.active
        have = {(ws.cell(r, 6).value, ws.cell(r, 2).value) for r in range(2, ws.max_row + 1)}
    else:
        wb = Workbook(); ws = wb.active; ws.title = '상승이유'
        ws.append(['날짜', '종목명', '이유', '기사 제목', '출처', '링크', '시각', '방향'])
        have = set()
    for x in sorted(add, key=lambda x: (x['d'], x['t'])):
        if (x['u'], x['n']) in have:
            continue
        ws.append([x['d'], x['n'], x['r'], x['h'], x['s'], x['u'], x['t'], x['dir']])
    wb.save(path)


def git(repo, *args):
    return subprocess.run(['git', '-C', repo, *args], capture_output=True, text=True, encoding='utf-8', errors='ignore')


def push(repo, today):
    if not os.path.isdir(os.path.join(repo, '.git')):
        log('repo 폴더가 git 저장소가 아니라서 업로드는 건너뜁니다. (GitHub 웹에서 data/reasons.json 을 직접 올려 주세요)')
        return
    r = git(repo, 'pull', '--rebase', '--autostash')
    if r.returncode: log('git pull 경고: ' + (r.stderr or r.stdout).strip()[:200])
    git(repo, 'add', 'data/reasons.json')
    if git(repo, 'diff', '--cached', '--quiet').returncode == 0:
        log('올릴 변경이 없습니다.'); return
    r = git(repo, 'commit', '-m', f'reasons {today:%Y-%m-%d}')
    if r.returncode: log('git commit 실패: ' + (r.stderr or r.stdout).strip()[:200]); return
    r = git(repo, 'push')
    log('GitHub 업로드 완료' if r.returncode == 0 else 'git push 실패: ' + (r.stderr or r.stdout).strip()[:300])


def main():
    cfg = load_cfg()
    now = datetime.datetime.now(KST)
    names = load_names(cfg['repo'])
    log(f'종목명 {len(names)}개 로드, 검색어 {cfg["queries"]}')
    items = fetch_items(cfg, now)
    log(f'기사 {len(items)}건 조회')
    rows = build_rows(items, names)
    jp = os.path.join(cfg['repo'], 'data', 'reasons.json')
    old = json.load(open(jp, encoding='utf-8')) if os.path.exists(jp) else []
    allr, add = merge(old, rows, cfg['days_keep'], now.date())
    save_json(jp, allr)
    save_excel(os.path.join(HERE, '상승이유_누적.xlsx'), add)
    log(f'종목이 확인된 기사 {len(rows)}건 중 새로 추가 {len(add)}건, 저장된 전체 {len(allr)}건')
    if cfg['push']:
        push(cfg['repo'], now)


if __name__ == '__main__':
    main()
